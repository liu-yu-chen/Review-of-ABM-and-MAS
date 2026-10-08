from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

root = Path(r"D:\Literature-Research-Agent")
doc = Document(root / "llm_current.docx")

def para_before(anchor, text, style="Normal"):
    p = OxmlElement("w:p")
    anchor._p.addprevious(p)
    out = Paragraph(p, anchor._parent)
    out.style = style
    if text:
        out.add_run(text)
    return out

def para_after(anchor, text, style="Normal"):
    p = OxmlElement("w:p")
    anchor._p.addnext(p)
    out = Paragraph(p, anchor._parent)
    out.style = style
    if text:
        out.add_run(text)
    return out

# Add a focused methods subsection immediately before the ABM trend results.
trend_heading = next(p for p in doc.paragraphs if p.text.strip() == "2. Research trends in agent-based modeling")
method_blocks = [
    ("1.1 DeepSeek API classification and prompt design", "Heading 2"),
    ("For corpus-level classification, eligible records with abstracts were sent to the DeepSeek API through an OpenAI-compatible asynchronous client, using the deepseek-chat model. Each request contained the paper title and abstract (abstracts were capped at 6,000 characters); records without an abstract were not sent for model classification. The request used temperature 0, JSON-object output, and a 512-token response limit. Concurrent requests were checkpointed and retryable failures were retried, allowing the run to resume without silently replacing completed records.", "Normal"),
    ("Prompt engineering was used to constrain the task rather than invite free-form interpretation. A fixed system prompt established the model’s role as an ABM and complex-systems literature classifier and supplied standardized taxonomies for primary topic, dominant simulated-agent type, LLM use, and LLM function. The model was instructed to use only evidence in the title and abstract, return exactly one topic and one dominant agent category, and avoid unsupported inference. An LLM was coded as used only when it contributed to the research method; background mentions and use limited to manuscript writing, editing, translation, or literature searching were excluded. Conventional machine-learning or NLP methods were not automatically treated as LLMs, and the LLM-agent category was reserved for cases where an LLM was part of the simulated agent architecture. The required output was a compact JSON object with fixed fields and taxonomy codes, which was parsed and checked before merging with the corpus. These constraints improve consistency and auditability, but do not remove classification uncertainty or the need for human validation.", "Normal"),
]
for text, style in method_blocks:
    para_before(trend_heading, text, style)

# Keep the pre/post subtheme comparison intact; revise only the LLM-focused
# findings and add the requested application-domain/agent-type synthesis.
section_heading = next(p for p in doc.paragraphs if p.text.strip() == "3. Before and after the emergence of large language models")
section_heading.text = "3. Large language model adoption in agent-based modeling"

heading = next(p for p in doc.paragraphs if p.text.strip() == "3.2 Adoption is accelerating, but the field is not yet LLM-dominated")
heading.text = "3.2 Diffusion is accelerating but remains at an early stage"

adoption = next(p for p in doc.paragraphs if p.text.startswith("Among 29,139 papers, 236 were identified as incorporating LLMs"))
adoption.text = (
    "LLM use is growing rapidly from a small base. Of the 29,139 papers in the through-2025 corpus, 236 (0.81%) were classified as using an LLM. Annual counts rose from 32 in 2023 to 63 in 2024 and 141 in 2025. Under the annual denominator used for the corpus trend analysis, the corresponding shares increased from 1.50% to 2.75% and 5.65% (a 2025-to-2023 rate ratio of 3.77). Thus, the signal is clear, but LLM use had not become standard practice by 2025: most ABM publications in the observation window were not classified as LLM-enabled. The short post-release window also makes these estimates an early-diffusion snapshot rather than evidence of a settled trajectory."
)

roles = next(p for p in doc.paragraphs if p.text.startswith("The distribution of LLM functions is even more informative than adoption alone."))
roles.paragraph_format.keep_together = True
roles.text = (
    "The function of the LLM matters as much as its presence. Among 235 papers with a classifiable role, 140 (59.6%; Wilson 95% CI 53.2–65.6) used the model within agent behavior or decision-making. The remaining 95 papers (40.4%) used LLMs in supporting functions, including code or model implementation, data and parameter generation, evaluation or text analysis, and natural-language interfaces. The corpus therefore points to two concurrent patterns: LLMs are often proposed as part of the simulated decision mechanism, while a substantial minority assist the research workflow around the model. These roles should be reported separately because only the former directly changes how simulated agents produce behavior."
)

interpretation = next(p for p in doc.paragraphs if p.text.startswith("This is the central transition identified by the review:"))
interpretation.text = (
    "The most consequential change is the introduction of language models into the agent’s decision process, rather than simply their use as a productivity tool. Conventional ABM typically specifies decision rules directly; LLM-enabled agents can interpret contextual information, draw on memory or retrieval, plan, and communicate in natural language (Park et al., 2023). This can represent flexible, language-mediated behavior that is difficult to capture with fixed rules. It also makes behavior dependent on the model, prompt, context, and stochastic settings, so plausibility alone cannot establish behavioral validity."
)

lim_heading = next(p for p in doc.paragraphs if p.text.strip() == "3.3 Why the LLM transition is not a simple replacement")
lim_heading.text = "3.3 Limits of current LLM-enabled ABM and research priorities"
lim_paras = [p for p in doc.paragraphs if p.text.startswith("The pre-LLM ABM tradition already contains mechanisms") or p.text.startswith("A useful distinction is between a generative agent")]
lim_paras[0].text = (
    "The corpus indicates an emerging approach, not a replacement of established ABM. LLM-enabled models still need to represent bounded rationality, adaptation, spatial and network constraints, and institutional rules, and they must be calibrated and validated against empirical evidence. A fluent or contextually plausible answer may nevertheless yield unstable choice probabilities, weak demographic differentiation, or unrealistic responses to policy. Prompt sensitivity, model updates, hallucinated content, computational cost, and limited reproducibility therefore constrain current applications (Ghaffarzadegan et al., 2024; Li et al., 2024). These concerns are especially consequential when an LLM is the decision engine, because changes in the model or prompt can alter the simulated mechanism itself."
)
lim_paras[1].text = (
    "Future studies should make this uncertainty measurable. At minimum, they should compare LLM agents with rule-based and empirically calibrated baselines; test sensitivity to model version, prompt, random seed, temperature, and information access; and validate both individual behavior and aggregate outcomes against observed or experimental benchmarks. Researchers should also report inference cost, data provenance, privacy safeguards, and the human review used to validate classifications or synthetic outputs. Because the observed LLM literature is concentrated in a few application domains and spans only the first years of diffusion, replication across domains and longer time series will be needed before claims of general-purpose value are warranted."
)

fig14 = next(p for p in doc.paragraphs if p.text.startswith("Figure 14. Cross-distribution of application domains and LLM roles"))
new_heading = para_after(fig14, "3.4 Application domains and simulated agent types", "Heading 2")
domain_paragraph = para_after(new_heading,
    "LLM-enabled ABM is not evenly distributed across application areas. Among the 235 classifiable papers shown in Figure 14, the largest groups concern social and opinion dynamics (81 papers; 34.5%) and health and epidemiology (58; 24.7%). They are followed by AI and multi-agent computing (28; 11.9%), policy and disaster response (25; 10.6%), sustainability and energy (21; 8.9%), and transportation and urban mobility (18; 7.7%); networks and communication and other specialized domains each contribute two papers. This pattern places current applications chiefly in social-behavioral and health-related problems, with a smaller but meaningful technical strand focused on AI and multi-agent systems. The cross-distribution also indicates that decision-making is the most common LLM function across the major domains, while coding, data generation, analysis, and interface roles provide supporting pathways.")
para_after(domain_paragraph,
    "The simulated entities in LLM-related studies also differ from the broader ABM corpus. The comparative analysis reports an 11.5-percentage-point higher representation of software or computational agents and a 12.7-point lower representation of biological entities in LLM papers. This is consistent with the emergence of software agents and LLM-based agents as explicit objects of study, but it does not mean that every paper using an LLM simulates an LLM agent: many use the model to support decisions, data preparation, analysis, or interaction around another agent population. Keeping the LLM’s functional role separate from the type of simulated agent is therefore essential for interpreting the field’s development.")

eval_heading = next(p for p in doc.paragraphs if p.text.strip() == "3.4 An evaluative framework for LLM-enabled ABM")
eval_heading.text = "3.5 An evaluative framework for LLM-enabled ABM"

out = root / "agent_based_modeling_review_revised_v3.docx"
doc.save(out)
print(out)
